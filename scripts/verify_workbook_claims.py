#!/usr/bin/env python3
import argparse
import json
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook


def as_decimal(value):
    if isinstance(value, bool) or value is None:
        return None
    try:
        return Decimal(str(value).replace(",", "").strip())
    except (InvalidOperation, ValueError):
        return None


def load_cell(source, base_dir: Path, cache: dict):
    workbook_path = Path(source["workbook"]).expanduser()
    if not workbook_path.is_absolute():
        workbook_path = (base_dir / workbook_path).resolve()
    if not workbook_path.is_file():
        raise ValueError(f"missing workbook: {workbook_path}")
    key = str(workbook_path)
    if key not in cache:
        cache[key] = (
            load_workbook(workbook_path, data_only=True, read_only=False),
            load_workbook(workbook_path, data_only=False, read_only=False),
        )
    data_book, formula_book = cache[key]
    sheet_name = source["sheet"]
    if sheet_name not in data_book.sheetnames:
        raise ValueError(f"missing sheet: {workbook_path.name} / {sheet_name}")
    coordinate = source["cell"]
    value = data_book[sheet_name][coordinate].value
    formula_value = formula_book[sheet_name][coordinate].value
    if value is None and isinstance(formula_value, str) and formula_value.startswith("="):
        raise ValueError(f"formula has no cached value: {workbook_path.name} / {sheet_name}!{coordinate}")
    return value, f"{workbook_path.name} / {sheet_name}!{coordinate}"


def verify_claim(claim, base_dir: Path, cache: dict):
    claim_id = claim.get("id", "unnamed")
    operation = claim.get("operation", "exact")
    sources = claim.get("sources") or []
    if operation not in {"exact", "sum"}:
        raise ValueError(f"{claim_id}: operation must be exact or sum")
    if not sources:
        raise ValueError(f"{claim_id}: sources are required")
    if operation == "exact" and len(sources) != 1:
        raise ValueError(f"{claim_id}: exact requires exactly one source")

    loaded = [load_cell(source, base_dir, cache) for source in sources]
    values = [item[0] for item in loaded]
    labels = [item[1] for item in loaded]
    expected = claim.get("expected")
    expected_number = as_decimal(expected)
    tolerance = as_decimal(claim.get("tolerance", 0)) or Decimal("0")

    if operation == "sum":
        numbers = [as_decimal(value) for value in values]
        if any(value is None for value in numbers):
            raise ValueError(f"{claim_id}: sum sources must all be numeric")
        actual = sum(numbers, Decimal("0"))
    else:
        actual = values[0]

    actual_number = as_decimal(actual)
    if expected_number is not None and actual_number is not None:
        passed = abs(actual_number - expected_number) <= tolerance
        display_actual = str(actual_number)
        display_expected = str(expected_number)
    else:
        passed = str(actual).strip() == str(expected).strip()
        display_actual = actual
        display_expected = expected

    return {
        "id": claim_id,
        "passed": passed,
        "expected": display_expected,
        "actual": display_actual,
        "sources": labels,
    }


def main() -> int:
    parser = argparse.ArgumentParser(description="Verify critical PPT facts against source workbook cells.")
    parser.add_argument("claims_file", help="JSON file containing claims")
    args = parser.parse_args()

    claims_path = Path(args.claims_file).expanduser().resolve()
    payload = json.loads(claims_path.read_text(encoding="utf-8"))
    claims = payload.get("claims")
    if not isinstance(claims, list) or not claims:
        parser.error("claims must be a non-empty list")

    cache = {}
    results = []
    errors = []
    for claim in claims:
        try:
            results.append(verify_claim(claim, claims_path.parent, cache))
        except Exception as exc:
            errors.append({"id": claim.get("id", "unnamed"), "error": str(exc)})

    failed = [item for item in results if not item["passed"]]
    report = {
        "status": "passed" if not failed and not errors else "blocked",
        "claims": len(claims),
        "passed": len(results) - len(failed),
        "failed": failed,
        "errors": errors,
        "results": results,
    }
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0 if report["status"] == "passed" else 1


if __name__ == "__main__":
    raise SystemExit(main())
