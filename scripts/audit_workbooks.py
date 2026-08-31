#!/usr/bin/env python3
import argparse
import datetime as dt
import hashlib
import json
from pathlib import Path

from openpyxl import load_workbook


STAR_MARKERS = ("*", "＊", "★", "☆")


def file_hash(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def json_value(value):
    if isinstance(value, (dt.datetime, dt.date, dt.time)):
        return value.isoformat()
    if isinstance(value, bytes):
        return value.hex()
    return value


def audit_workbook(path: Path) -> dict:
    keep_vba = path.suffix.lower() == ".xlsm"
    formulas = load_workbook(path, data_only=False, read_only=False, keep_vba=keep_vba)
    values = load_workbook(path, data_only=True, read_only=False, keep_vba=keep_vba)
    sheets = []

    for sheet in formulas.worksheets:
        value_sheet = values[sheet.title]
        rows = []
        starred_cells = []
        nonempty_count = 0

        for row in sheet.iter_rows():
            row_cells = []
            for cell in row:
                raw = cell.value
                if raw is None:
                    continue
                nonempty_count += 1
                is_formula = cell.data_type == "f" or (isinstance(raw, str) and raw.startswith("="))
                cached = value_sheet[cell.coordinate].value if is_formula else None
                is_starred = isinstance(raw, str) and any(marker in raw for marker in STAR_MARKERS)
                record = {
                    "cell": cell.coordinate,
                    "value": json_value(raw),
                    "value_type": cell.data_type,
                    "number_format": cell.number_format,
                }
                if is_formula:
                    record["formula"] = raw
                    record["cached_value"] = json_value(cached)
                if is_starred:
                    record["starred"] = True
                    starred_cells.append(cell.coordinate)
                row_cells.append(record)
            if row_cells:
                rows.append({"row": row[0].row, "cells": row_cells})

        sheets.append(
            {
                "name": sheet.title,
                "state": sheet.sheet_state,
                "max_row": sheet.max_row,
                "max_column": sheet.max_column,
                "nonempty_cells": nonempty_count,
                "merged_ranges": [str(item) for item in sheet.merged_cells.ranges],
                "starred_cells": starred_cells,
                "rows": rows,
            }
        )

    return {
        "path": str(path.resolve()),
        "filename": path.name,
        "size_bytes": path.stat().st_size,
        "sha256": file_hash(path),
        "sheet_count": len(sheets),
        "sheet_names": [sheet["name"] for sheet in sheets],
        "sheets": sheets,
    }


def main() -> int:
    parser = argparse.ArgumentParser(description="Audit every sheet and non-empty cell in XLSX/XLSM sources.")
    parser.add_argument("workbooks", nargs="+", help="Source .xlsx or .xlsm files")
    parser.add_argument("--output", required=True, help="Output JSON ledger")
    args = parser.parse_args()

    paths = [Path(item).expanduser() for item in args.workbooks]
    invalid = [str(path) for path in paths if not path.is_file() or path.suffix.lower() not in {".xlsx", ".xlsm"}]
    if invalid:
        parser.error(f"Unsupported or missing workbook(s): {', '.join(invalid)}")

    result = {
        "generated_at": dt.datetime.now(dt.timezone.utc).isoformat(),
        "workbook_count": len(paths),
        "workbooks": [audit_workbook(path) for path in paths],
    }
    output = Path(args.output).expanduser()
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({
        "output": str(output.resolve()),
        "workbooks": len(result["workbooks"]),
        "sheets": sum(item["sheet_count"] for item in result["workbooks"]),
        "starred_cells": sum(len(sheet["starred_cells"]) for item in result["workbooks"] for sheet in item["sheets"]),
    }, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
